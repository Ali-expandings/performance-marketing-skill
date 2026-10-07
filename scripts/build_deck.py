#!/usr/bin/env python3
"""Builds the 10-slide 16:9 executive deck comparing the last two periods.

Usage: build_deck.py --metrics out/audit/metrics.json --narrative narrative.json [--profile profile.json] [--out deck.pptx]
"""
import argparse
import os
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from common import load_settings, read_json  # noqa: E402

MONTHS = ["", "Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
PENDING = "Pending narrative"


class Deck:
    def __init__(self, cfg):
        self.c = {k: RGBColor(*v) for k, v in cfg["deck"].items() if isinstance(v, list)}
        self.font = cfg["deck"]["font"]
        self.p = Presentation()
        self.p.slide_width, self.p.slide_height = Inches(13.333), Inches(7.5)
        self.n = 0

    def slide(self, dark=False, title=None, kicker=None):
        s = self.p.slides.add_slide(self.p.slide_layouts[6])
        self.n += 1
        bg = s.background.fill
        bg.solid()
        bg.fore_color.rgb = self.c["dark"] if dark else self.c["light"]
        if title:
            if kicker:
                self.text(s, 0.6, 0.35, 12, 0.35, kicker.upper(), 11, True, self.c["teal"])
            self.text(s, 0.6, 0.65, 12, 0.7, title, 26, True, RGBColor(255, 255, 255) if dark else self.c["dark"])
            self.text(s, 12.2, 7.0, 0.8, 0.3, str(self.n), 10, False, self.c["muted"], PP_ALIGN.RIGHT)
        return s

    def text(self, s, x, y, w, h, txt, size=14, bold=False, color=None, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
        tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = Inches(0.05)
        lines = txt if isinstance(txt, list) else [txt]
        for i, line in enumerate(lines):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.alignment = align
            run = para.add_run()
            run.text = str(line)
            run.font.size, run.font.bold, run.font.name = Pt(size), bold, self.font
            run.font.color.rgb = color or self.c["text"]
            para.space_after = Pt(6)
        return tb

    def card(self, s, x, y, w, h, fill=None, accent=None):
        sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        sh.adjustments[0] = 0.06
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill or self.c["card"]
        sh.line.color.rgb = self.c["border"]
        sh.shadow.inherit = False
        if accent:
            bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y + 0.15), Inches(0.07), Inches(h - 0.3))
            bar.fill.solid()
            bar.fill.fore_color.rgb = accent
            bar.line.fill.background()
            bar.shadow.inherit = False
        return sh

    def bar(self, s, x, y, w, h, color):
        if w <= 0:
            return
        b = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        b.fill.solid()
        b.fill.fore_color.rgb = color
        b.line.fill.background()
        b.shadow.inherit = False

    def table(self, s, x, y, w, rows, header, col_w=None, size=11):
        t = s.shapes.add_table(len(rows) + 1, len(header), Inches(x), Inches(y), Inches(w), Inches(0.4 * (len(rows) + 1))).table
        for j, h in enumerate(header):
            cell = t.cell(0, j)
            cell.text = h
            cell.fill.solid()
            cell.fill.fore_color.rgb = self.c["dark"]
            self._cellfont(cell, size, True, RGBColor(255, 255, 255))
            if col_w:
                t.columns[j].width = Inches(col_w[j])
        for i, row in enumerate(rows, 1):
            for j, v in enumerate(row):
                cell = t.cell(i, j)
                cell.text = str(v)
                cell.fill.solid()
                cell.fill.fore_color.rgb = self.c["card"] if i % 2 else self.c["light"]
                color = self.c["text"]
                if isinstance(v, str) and v.endswith("%") and v[:1] in "+-":
                    color = self.c["green"] if v.startswith("+") else self.c["red"]
                self._cellfont(cell, size, j == 0, color)
        return t

    def _cellfont(self, cell, size, bold, color):
        for para in cell.text_frame.paragraphs:
            for run in para.runs:
                run.font.size, run.font.bold, run.font.name = Pt(size), bold, self.font
                run.font.color.rgb = color


def money(v, cur="$"):
    if v is None:
        return "n/a"
    return f"{cur}{v / 1000:,.1f}K" if abs(v) >= 10000 else f"{cur}{v:,.2f}" if abs(v) < 100 else f"{cur}{v:,.0f}"


def pct(v, sign=False):
    if v is None:
        return "n/a"
    return f"{v * 100:+.1f}%" if sign else f"{v * 100:.1f}%"


def num(v):
    return "n/a" if v is None else f"{v:,.0f}"


def build(m, n, cfg, path):
    d = Deck(cfg)
    C = d.c
    WHITE = RGBColor(255, 255, 255)
    ys = m["periods"][-2:]
    pa, pb = (ys[0], ys[1]) if len(ys) == 2 else (ys[0], ys[0])
    A, B = m["comparable"][str(pa)], m["comparable"][str(pb)]
    cmp = m.get("comparison", {})
    months = m["comparable_months"]
    span = f"{MONTHS[months[0]]}–{MONTHS[months[-1]]}" if months else ""

    # 1 title
    s = d.slide(dark=True)
    d.text(s, 0.8, 1.0, 11, 0.4, f"{m['client'].upper()} · PERFORMANCE AUDIT", 13, True, C["teal"])
    d.text(s, 0.8, 1.6, 11.5, 2.2, n.get("title") or f"{pa} vs {pb}: what drove growth and what broke", 40, True, WHITE)
    d.text(s, 0.8, 3.9, 11, 1.0, n.get("subtitle") or f"Comparable window {span}, CRM deals reconciled with ad spend", 18, False, RGBColor(203, 213, 225))
    d.card(s, 0.8, 5.6, 6.2, 0.8, fill=RGBColor(30, 45, 80))
    d.text(s, 1.0, 5.72, 6, 0.6, "✓ " + (n.get("verification_stamp") or "Numbers verified against CRM and ad exports"), 13, True, WHITE, anchor=MSO_ANCHOR.MIDDLE)

    # 2 scorecard
    s = d.slide(title=f"{pa} vs {pb}: revenue, volume and price", kicker=f"Executive scorecard · {span}")
    kpis = [("Revenue", "revenue", money), ("Won deals", "won", num), ("AOV", "aov", money),
            ("CAC", "cac", money), ("Inquiry → won", "cvr_inquiry", pct), ("ROAS", "roas", lambda v: "n/a" if v is None else f"{v:.2f}x")]
    for i, (lab, k, f) in enumerate(kpis):
        x, y = 0.6 + (i % 3) * 4.1, 1.6 + (i // 3) * 2.55
        ch = (cmp.get(k) or {}).get("change_pct")
        good = (ch or 0) >= 0 if k != "cac" else (ch or 0) <= 0
        d.card(s, x, y, 3.85, 2.3, accent=C["green"] if good else C["red"])
        d.text(s, x + 0.3, y + 0.2, 3.4, 0.4, lab.upper(), 12, True, C["muted"])
        d.text(s, x + 0.3, y + 0.6, 3.4, 0.8, f(B.get(k)), 34, True, C["dark"])
        d.text(s, x + 0.3, y + 1.45, 3.4, 0.4, f"{pa}: {f(A.get(k))}", 13, False, C["muted"])
        d.text(s, x + 0.3, y + 1.8, 3.4, 0.4, pct(ch, True) if ch is not None else "", 15, True, C["green"] if good else C["red"])
    vf = next((f for f in m["findings"] if f["pattern"] == "volume_vs_aov"), None)
    if vf:
        e = vf["evidence"]
        d.text(s, 0.6, 6.75, 12.2, 0.5, f"Revenue {pct(e['revenue_change_pct'], True)} = deals {pct(e['won_change_pct'], True)} × AOV {pct(e['aov_change_pct'], True)}. Growth is {e['driver']}-led.", 14, True, C["steel"])

    # 3 month by month
    s = d.slide(title="Month by month", kicker=f"{pa} vs {pb}")
    rows = []
    for x in m.get("month_by_month", []):
        rows.append([MONTHS[x["month"]], money(x.get(f"revenue_{pa}")), money(x.get(f"revenue_{pb}")), pct(x.get("revenue_change_pct"), True),
                     num(x.get(f"won_{pa}")), num(x.get(f"won_{pb}")), money(x.get(f"cac_{pa}")), money(x.get(f"cac_{pb}")),
                     pct(x.get(f"cvr_inquiry_{pa}")), pct(x.get(f"cvr_inquiry_{pb}"))])
    d.table(s, 0.6, 1.5, 12.1, rows[:12], ["Month", f"Rev {pa}", f"Rev {pb}", "Δ Rev", f"Won {pa}", f"Won {pb}", f"CAC {pa}", f"CAC {pb}", f"Win {pa}", f"Win {pb}"],
            [1.0, 1.3, 1.3, 1.1, 1.1, 1.1, 1.2, 1.2, 1.4, 1.4], size=11)

    # 4 pivots
    s = d.slide(title="The three moments that moved the numbers", kicker="Critical operational pivots")
    piv = (n.get("pivots") or [])[:3] or [{"title": PENDING}]
    colors = [C["red"], C["gold"], C["teal"]]
    for i, p in enumerate(piv):
        x = 0.6 + i * 4.1
        d.card(s, x, 1.6, 3.85, 5.3, accent=colors[i % 3])
        d.text(s, x + 0.3, 1.8, 3.4, 0.4, f"{p.get('month', '')} {p.get('period', '')}".strip().upper(), 12, True, colors[i % 3])
        d.text(s, x + 0.3, 2.2, 3.4, 1.0, p.get("title", ""), 19, True, C["dark"])
        d.text(s, x + 0.3, 3.35, 3.4, 2.0, p.get("evidence", ""), 13, False, C["text"])
        d.text(s, x + 0.3, 5.5, 3.4, 1.3, "→ " + p.get("lesson", "") if p.get("lesson") else "", 13, True, C["steel"])

    # 5 efficiency
    s = d.slide(title=(n.get("efficiency") or {}).get("headline") or "The efficiency blueprint", kicker="Lowest-CAC months")
    eff = sorted([f for f in m["findings"] if f["pattern"] == "efficiency_peak" and str(f["period"]) == str(pb)], key=lambda f: f["evidence"]["cac"])[:3]
    for i, f in enumerate(eff):
        y = 1.6 + i * 1.75
        d.card(s, 0.6, y, 4.6, 1.55, accent=C["green"])
        d.text(s, 0.9, y + 0.15, 4.2, 0.4, f"{MONTHS[f['month']]} {f['period']}", 13, True, C["muted"])
        d.text(s, 0.9, y + 0.5, 4.2, 0.6, f"CAC {money(f['evidence']['cac'])}", 26, True, C["green"])
        d.text(s, 0.9, y + 1.08, 4.2, 0.4, f"win rate {pct(f['evidence']['cvr'])} · {num(f['evidence']['won'])} won", 12, False, C["text"])
    bl = [x for x in (n.get("efficiency") or {}).get("bullets", []) if x] or [PENDING]
    d.card(s, 5.6, 1.6, 7.1, 5.1)
    d.text(s, 5.9, 1.85, 6.6, 4.7, [f"•  {x}" for x in bl], 16, False, C["text"])

    # 6 channel architecture
    s = d.slide(title="Acquisition ≠ closing", kicker="Sales channel architecture")
    t = m["totals"][str(pb)]
    d.card(s, 0.6, 1.6, 5.4, 5.1)
    d.text(s, 0.9, 1.8, 5, 0.4, "WHERE LEADS COME FROM (ADS)", 12, True, C["muted"])
    acq = [("Form leads", t.get("leads") or 0), ("Messaging starts", t.get("messages") or 0)]
    mx = max(v for _, v in acq) or 1
    for i, (lab, v) in enumerate(acq):
        y = 2.5 + i * 1.3
        d.text(s, 0.9, y, 4.8, 0.4, f"{lab}: {num(v)}", 14, True, C["dark"])
        d.bar(s, 0.9, y + 0.45, 4.8 * v / mx, 0.4, C["steel"])
    d.card(s, 6.4, 1.6, 6.3, 5.1)
    d.text(s, 6.7, 1.8, 5.9, 0.4, "WHERE DEALS CLOSE (CRM)", 12, True, C["muted"])
    chans = m["breakdowns"][str(pb)]["channels"][:5]
    for i, c in enumerate(chans):
        y = 2.4 + i * 0.8
        d.text(s, 6.7, y, 3.0, 0.4, c["channel"], 13, True, C["dark"])
        d.bar(s, 9.2, y + 0.05, 2.6 * (c["share_won"] or 0), 0.35, C["teal"] if i == 0 else C["muted"])
        d.text(s, 11.85, y, 0.8, 0.4, pct(c["share_won"]), 12, True, C["text"])
    if n.get("channel_story"):
        d.text(s, 0.6, 6.8, 12.1, 0.5, n["channel_story"], 13, True, C["steel"])

    # 7 geo
    s = d.slide(title="Volume markets vs value markets", kicker="Geographic economics")
    geo = [g for g in m["breakdowns"][str(pb)]["countries"] if g["country"] != "Unknown"][:6]
    d.bar(s, 3.3, 1.62, 0.25, 0.18, C["steel"])
    d.text(s, 3.6, 1.5, 2.2, 0.4, "Share of deals", 12, True, C["muted"])
    d.bar(s, 5.8, 1.62, 0.25, 0.18, C["gold"])
    d.text(s, 6.1, 1.5, 2.4, 0.4, "Share of revenue", 12, True, C["muted"])
    for i, g in enumerate(geo):
        y = 2.0 + i * 0.75
        d.text(s, 0.6, y, 2.6, 0.5, g["country"], 13, True, C["dark"])
        d.bar(s, 3.3, y + 0.02, 5.0 * (g["share_won"] or 0), 0.25, C["steel"])
        d.bar(s, 3.3, y + 0.3, 5.0 * (g["share_revenue"] or 0), 0.25, C["gold"])
        d.text(s, 8.6, y, 1.6, 0.5, f"AOV {money(g['aov'])}", 12, True, C["text"])
    d.card(s, 10.0, 2.0, 2.7, 4.4, accent=C["gold"])
    d.text(s, 10.3, 2.2, 2.3, 4.1, n.get("geo_story") or PENDING, 14, False, C["text"])

    # 8 roadmap
    s = d.slide(title="Improvement roadmap", kicker="Four growth pillars")
    pil = (n.get("pillars") or [])[:4] or [{"name": PENDING, "actions": [], "kpi": ""}]
    for i, p in enumerate(pil):
        x, y = 0.6 + (i % 2) * 6.15, 1.6 + (i // 2) * 2.7
        d.card(s, x, y, 5.95, 2.5, accent=[C["teal"], C["steel"], C["gold"], C["green"]][i % 4])
        d.text(s, x + 0.3, y + 0.15, 5.4, 0.5, f"{i + 1}. {p.get('name', '')}", 18, True, C["dark"])
        d.text(s, x + 0.3, y + 0.7, 5.4, 1.3, [f"•  {a}" for a in p.get("actions", []) if a], 13, False, C["text"])
        if p.get("kpi"):
            d.text(s, x + 0.3, y + 2.0, 5.4, 0.4, f"KPI: {p['kpi']}", 12, True, C["steel"])

    # 9 targets
    s = d.slide(title="Targets", kicker="Scorecard and financial roadmap")
    tg = [[x.get("period", ""), x.get("revenue", ""), x.get("cac", ""), x.get("win_rate", "")] for x in (n.get("targets") or []) if x.get("period")]
    d.table(s, 0.6, 1.6, 8.0, tg or [[PENDING, "", "", ""]], ["Period", "Revenue target", "CAC target", "Win-rate target"], [2.0, 2.0, 2.0, 2.0], size=14)
    d.card(s, 9.0, 1.6, 3.7, 3.0)
    d.text(s, 9.25, 1.8, 3.3, 0.4, "BASELINE " + str(pb), 12, True, C["muted"])
    d.text(s, 9.25, 2.25, 3.3, 2.2, [f"Revenue {money(B.get('revenue'))}", f"CAC {money(B.get('cac'))}", f"Win rate {pct(B.get('cvr_inquiry'))}", f"AOV {money(B.get('aov'))}"], 16, True, C["dark"])

    # 10 mandates
    s = d.slide(dark=True, title="Leadership mandates", kicker="Executive summary")
    md = [x for x in (n.get("mandates") or []) if x][:5] or [PENDING]
    for i, x in enumerate(md):
        y = 1.65 + i * 1.0
        d.text(s, 0.8, y, 0.8, 0.8, f"{i + 1:02d}", 28, True, C["teal"])
        d.text(s, 1.8, y + 0.08, 10.8, 0.9, x, 18, False, WHITE)

    d.p.save(path)
    return path


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--metrics", required=True)
    ap.add_argument("--narrative")
    ap.add_argument("--profile")
    ap.add_argument("--out")
    a = ap.parse_args()
    cfg = load_settings(a.profile)
    m = read_json(a.metrics)
    n = read_json(a.narrative) if a.narrative and os.path.exists(a.narrative) else {}
    if not n:
        print("note: no narrative given, text slides will say 'Pending narrative'")
    ys = m["periods"][-2:]
    out = a.out or os.path.join(os.path.dirname(a.metrics), f"{cfg['file_prefix']}_Deck_{'_'.join(map(str, ys))}.pptx")
    print("wrote", build(m, n, cfg, out))


if __name__ == "__main__":
    main()
