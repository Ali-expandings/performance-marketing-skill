#!/usr/bin/env python3
"""Grades one audit run out of 100 and logs the score so runs can be compared over time.

Checks: numbers in the narrative trace back to metrics (40), every high finding is addressed (20),
narrative is complete (15), data caveats are acknowledged (10), workbook and deck are valid (15).

Usage: selfcheck.py --metrics out/audit/metrics.json --narrative narrative.json [--workbook X.xlsx ...] [--deck X.pptx]
"""
import argparse
import datetime
import json
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from common import SKILL_DIR, allowed_values, collect_numbers, is_verified, load_settings, numbers_in, read_json  # noqa: E402

MONTHS = ["", "jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"]
EXEMPT_KEYS = {"targets", "_instructions", "_high_findings_to_address", "period", "month", "quarter"}


def walk_text(obj, path=""):
    if isinstance(obj, dict):
        for k, v in obj.items():
            if k in EXEMPT_KEYS:
                continue
            yield from walk_text(v, f"{path}.{k}")
    elif isinstance(obj, list):
        for i, v in enumerate(obj):
            yield from walk_text(v, f"{path}[{i}]")
    elif isinstance(obj, str):
        yield path, obj


def check_numbers(m, n, cfg):
    allowed = allowed_values(collect_numbers(m) + collect_numbers(cfg["detectors"]))
    bad, total = [], 0
    for path, text in walk_text(n):
        for v, raw, tol in numbers_in(text):
            if 1900 <= v <= 2100 or (v.is_integer() and 0 <= v <= 12 and "%" not in raw and "$" not in raw):
                continue  # years, small counts like "3 archetypes"
            total += 1
            if not is_verified(v, allowed, tol):
                bad.append(f"{path}: '{raw}'")
    return total, bad


def addressed(f, n):
    text = json.dumps(n, ensure_ascii=False).lower()
    if f.get("month"):
        mon = MONTHS[int(f["month"])]
        return mon in text or f"-{int(f['month']):02d}" in text
    key = {"owner_history_guard": ["owner", "start date", "hire"], "closing_concentration": ["backup", "single point"],
           "data_quality": ["caveat"], "ad_concentration": ["creative pipeline", "one ad"]}.get(f["pattern"], [f["pattern"].replace("_", " ")])
    return any(k in text for k in key)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--metrics", required=True)
    ap.add_argument("--narrative", required=True)
    ap.add_argument("--workbook", nargs="*", default=[])
    ap.add_argument("--deck")
    ap.add_argument("--profile")
    ap.add_argument("--label", default="")
    ap.add_argument("--no-log", action="store_true", help="do not record this run (used by selftest)")
    a = ap.parse_args()
    cfg = load_settings(a.profile)
    m, n = read_json(a.metrics), read_json(a.narrative)
    issues, score = [], 0

    total, bad = check_numbers(m, n, cfg)
    s1 = 40 if total == 0 else round(40 * (1 - len(bad) / total))
    score += s1
    issues += [f"UNVERIFIED NUMBER {b} (not found in metrics.json: fix or remove)" for b in bad]

    highs = [f for f in m["findings"] if f["severity"] == "high"]
    miss = [f for f in highs if not addressed(f, n)]
    s2 = 20 if not highs else round(20 * (1 - len(miss) / len(highs)))
    score += s2
    issues += [f"HIGH FINDING NOT ADDRESSED: {f['pattern']} {f.get('period')}-{f.get('month')}" for f in miss]

    req = {"title": 1, "pivots": 3, "pillars": 4, "mandates": 5, "evolution": 1, "targets": 1}
    gaps = []
    for k, cnt in req.items():
        v = n.get(k)
        filled = [x for x in (v if isinstance(v, list) else [v]) if x and (not isinstance(x, dict) or any(x.values()))]
        if len(filled) < cnt:
            gaps.append(f"{k} needs {cnt}, has {len(filled)}")
    if not (n.get("efficiency") or {}).get("headline"):
        gaps.append("efficiency.headline empty")
    s3 = max(0, 15 - 3 * len(gaps))
    score += s3
    issues += [f"INCOMPLETE: {g}" for g in gaps]

    dqs = [f for f in m["findings"] if f["pattern"] == "data_quality"]
    cav = " ".join(n.get("caveats") or []).lower()
    s4 = 10 if not dqs or (cav and len([c for c in n.get("caveats", []) if c]) >= min(2, len(dqs))) else (5 if cav else 0)
    score += s4
    if s4 < 10:
        issues.append(f"CAVEATS: {len(dqs)} data-quality findings but caveats are thin. List each one.")

    s5 = 15
    for wb in a.workbook:
        try:
            from openpyxl import load_workbook
            sheets = load_workbook(wb, read_only=True).sheetnames
            if len(sheets) != 12:
                s5 -= 5
                issues.append(f"WORKBOOK {wb}: {len(sheets)} tabs, expected 12")
        except Exception as e:  # noqa: BLE001
            s5 -= 8
            issues.append(f"WORKBOOK {wb} unreadable: {e}")
    if a.deck:
        try:
            from pptx import Presentation
            p = Presentation(a.deck)
            if len(p.slides) != 10:
                s5 -= 5
                issues.append(f"DECK: {len(p.slides)} slides, expected 10")
            pend = sum("Pending narrative" in sh.text_frame.text for sl in p.slides for sh in sl.shapes if sh.has_text_frame)
            if pend:
                s5 -= 5
                issues.append(f"DECK: {pend} placeholders still say 'Pending narrative'")
        except Exception as e:  # noqa: BLE001
            s5 -= 8
            issues.append(f"DECK unreadable: {e}")
    if not a.workbook and not a.deck:
        s5 = 0
        issues.append("ARTIFACTS: no workbook or deck passed to the check")
    score += max(0, s5)

    pass_score = cfg["selfcheck"]["pass_score"]
    if bad:  # one invented number is enough to fail the run
        score = min(score, pass_score - 1)
    out = os.path.join(os.path.dirname(a.metrics), "selfcheck.md")
    lines = [f"# Self-check: {score}/100 ({'PASS' if score >= pass_score else 'FIX AND RERUN'}, pass mark {pass_score})", "",
             f"- numbers verified: {s1}/40 ({total - len(bad)}/{total})", f"- high findings addressed: {s2}/20 ({len(highs) - len(miss)}/{len(highs)})",
             f"- completeness: {s3}/15", f"- caveats: {s4}/10", f"- artifacts: {max(0, s5)}/15", "", "## Issues"] + [f"- {i}" for i in issues or ["none"]]

    os.makedirs(os.path.join(SKILL_DIR, "learnings"), exist_ok=True)
    log = os.path.join(SKILL_DIR, "learnings", "runs.jsonl")
    prev = []
    if os.path.exists(log):
        prev = [json.loads(x) for x in open(log, encoding="utf-8") if x.strip()]
    entry = {"date": datetime.date.today().isoformat(), "label": a.label, "score": score,
             "parts": {"numbers": s1, "findings": s2, "complete": s3, "caveats": s4, "artifacts": max(0, s5)},
             "patterns": sorted({f["pattern"] for f in m["findings"]}), "issue_types": sorted({i.split(":")[0] for i in issues})}
    if not a.no_log:
        with open(log, "a", encoding="utf-8") as f:
            f.write(json.dumps(entry) + "\n")
    if prev:
        avg = sum(p["score"] for p in prev[-5:]) / len(prev[-5:])
        lines += ["", f"Trend: this run {score} vs average {avg:.0f} of the last {len(prev[-5:])} runs."]
        recurring = {}
        for p in prev[-10:] + [entry]:
            for t in p["issue_types"]:
                recurring[t] = recurring.get(t, 0) + 1
        rec = [t for t, c in recurring.items() if c >= 3]
        if rec:
            lines += [f"Recurring issue types (3+ runs): {rec}. Log a lesson with learn.py and propose a rule."]
    open(out, "w", encoding="utf-8").write("\n".join(lines) + "\n")
    print("\n".join(lines))
    sys.exit(0 if score >= pass_score else 1)


if __name__ == "__main__":
    main()
